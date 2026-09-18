"""Render the two plate-reading plan Markdown files as .docx.

Written for: whoever needs the plans in a format the department circulates.
The Markdown files remain the source of truth; this script only re-renders them,
so edit the .md and run this again rather than editing the .docx.

    python docs/_build_plate_plans_docx.py

Supports the Markdown subset the plans actually use: ATX headings, paragraphs,
bullet and ordered lists with one level of nesting, GFM pipe tables (with
escaped \\| inside cells), fenced code blocks, horizontal rules, and the inline
run of `code`, **bold**, *italic* and [text](target).
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent

# Same palette as final_report/_build_hld.py so the set reads as one family.
NAVY = RGBColor(0x13, 0x2B, 0x4C)
NAVY_2 = RGBColor(0x1C, 0x41, 0x73)
INK = RGBColor(0x0F, 0x1A, 0x2B)
BODY = RGBColor(0x2E, 0x39, 0x48)
MUTED = RGBColor(0x5C, 0x69, 0x7C)
CODE = RGBColor(0x8A, 0x2B, 0x2B)

SANS = "Segoe UI"
MONO = "Consolas"

DOCS = [
    ("PLATE_READING_SHORT_TERM_PLAN.md", "Plate_Reading_Short_Term_Plan.docx"),
    ("PLATE_READING_LONG_TERM_PLAN.md", "Plate_Reading_Long_Term_Plan.docx"),
]


# ---------------------------------------------------------------- docx helpers
def shade(cell, hex_colour):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def borders(table, colour="C9D2DF"):
    tbl = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), colour)
        b.append(e)
    tbl.append(b)


def repeat_header(row):
    tr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr.append(el)


# --------------------------------------------------------------- inline markup
INLINE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<link>\[[^\]]+\]\([^)]+\))"
    r"|(?P<bold>\*\*[^*]+\*\*)"
    r"|(?P<italic>(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))"
)


def emit_inline(paragraph, text, *, size=10.5, colour=BODY, bold=False):
    """Append `text` to `paragraph`, honouring the inline Markdown subset."""
    text = text.replace(r"\|", "|")
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            add_run(paragraph, text[pos:m.start()], size=size, colour=colour, bold=bold)
        kind = m.lastgroup
        raw = m.group()
        if kind == "code":
            add_run(paragraph, raw[1:-1], size=size - 0.5, colour=CODE, bold=bold, mono=True)
        elif kind == "link":
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", raw).groups()
            add_run(paragraph, label, size=size, colour=colour, bold=bold)
            add_run(paragraph, f" ({target})", size=size - 1, colour=MUTED)
        elif kind == "bold":
            add_run(paragraph, raw[2:-2], size=size, colour=colour, bold=True)
        else:
            add_run(paragraph, raw[1:-1], size=size, colour=colour, bold=bold, italic=True)
        pos = m.end()
    if pos < len(text):
        add_run(paragraph, text[pos:], size=size, colour=colour, bold=bold)


def add_run(paragraph, text, *, size=10.5, colour=BODY, bold=False, italic=False, mono=False):
    if not text:
        return None
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    run.font.name = MONO if mono else SANS
    if mono:  # python-docx sets only the latin font; pin the others too
        rpr = run._element.get_or_add_rPr().rFonts
        for attr in ("w:eastAsia", "w:cs", "w:hAnsi"):
            rpr.set(qn(attr), MONO)
    return run


# ------------------------------------------------------------------- renderers
def heading(doc, text, level):
    sizes = {1: 15, 2: 12.5, 3: 11, 4: 10.5}
    colours = {1: NAVY, 2: NAVY_2, 3: NAVY_2, 4: INK}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level <= 1 else 11)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    emit_inline(p, text, size=sizes.get(level, 10.5), colour=colours.get(level, INK), bold=True)
    return p


def paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    emit_inline(p, text)
    return p


def list_item(doc, text, *, ordered, level):
    style = "List Number" if ordered and level == 0 else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.28 + 0.26 * level)
    if ordered and level > 0:  # Word restarts nested numbering unreliably; keep the literal marker
        emit_inline(p, text)
    else:
        emit_inline(p, text)
    return p


def code_block(doc, lines, language=""):
    table = doc.add_table(rows=1, cols=1)
    borders(table, colour="D8DEE8")
    cell = table.cell(0, 0)
    cell.text = ""
    shade(cell, "F5F7FB")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(1 if i == 0 else 0)
        add_run(p, line or " ", size=8.5, colour=INK, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "C9D2DF")
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def split_row(line):
    """Split a GFM table row, honouring \\| escapes."""
    body = line.strip().strip("|")
    cells = re.split(r"(?<!\\)\|", body)
    return [c.strip() for c in cells]


def table_block(doc, rows, usable_width):
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    size = 9.5 if cols <= 4 else (8.5 if cols <= 6 else 7.5)

    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)

    # Column widths proportional to content, clamped so no column collapses.
    weights = []
    for c in range(cols):
        longest = max(len(re.sub(r"[`*]", "", r[c])) for r in rows)
        weights.append(max(6, min(longest, 46)))
    total = sum(weights)
    widths = [max(0.6, usable_width * w / total) for w in weights]
    scale = usable_width / sum(widths)
    widths = [w * scale for w in widths]

    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            tc = table.cell(r, c)
            tc.text = ""
            tc.width = Inches(widths[c])
            p = tc.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if r == 0:
                emit_inline(p, text, size=size, colour=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                emit_inline(p, text, size=size)
            shade(tc, "132B4C" if r == 0 else ("FFFFFF" if r % 2 else "F5F7FB"))
    repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ----------------------------------------------------------------------- parse
LIST_RE = re.compile(r"^(?P<indent> *)(?P<marker>[-*+]|\d+\.)\s+(?P<text>.*)$")


def render(md_path: Path, out_path: Path) -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Inches(0.7))
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    usable = 8.27 - 1.4

    normal = doc.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # fenced code
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            # trim a common indent so nested blocks are not pushed off the page
            indents = [len(b) - len(b.lstrip()) for b in block if b.strip()]
            cut = min(indents) if indents else 0
            code_block(doc, [b[cut:] for b in block], language)
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]
        ):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            table_block(doc, rows, usable)
            continue

        # horizontal rule
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            rule(doc)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not title_done:
                title_done = True
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                add_run(p, text, size=20, colour=NAVY, bold=True)
                sub = doc.add_paragraph()
                sub.paragraph_format.space_after = Pt(10)
                add_run(
                    sub,
                    f"Gujarat Police CCTV platform · source: docs/{md_path.name} · "
                    f"rendered {date.today().isoformat()}",
                    size=9,
                    colour=MUTED,
                    italic=True,
                )
                rule(doc)
            else:
                heading(doc, text, level)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(6)
            emit_inline(p, stripped.lstrip("> ").strip(), colour=MUTED)
            i += 1
            continue

        # list item (with continuation lines)
        m = LIST_RE.match(line)
        if m:
            indent = len(m.group("indent"))
            level = 0 if indent < 2 else (1 if indent < 5 else 2)
            ordered = m.group("marker").endswith(".")
            text = m.group("text").strip()
            marker = m.group("marker")
            i += 1
            while (
                i < len(lines)
                and lines[i].strip()
                and not LIST_RE.match(lines[i])
                and not lines[i].strip().startswith(("|", "#", "```", "---"))
                and len(lines[i]) - len(lines[i].lstrip()) > indent
            ):
                text += " " + lines[i].strip()
                i += 1
            if ordered and level > 0:
                text = f"{marker} {text}"
            list_item(doc, text, ordered=ordered and level == 0, level=level)
            continue

        # paragraph (join soft-wrapped lines)
        block = [stripped]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].strip().startswith(("|", "#", "```", "---", ">"))
            and not LIST_RE.match(lines[i])
        ):
            block.append(lines[i].strip())
            i += 1
        paragraph(doc, " ".join(block))

    doc.save(out_path)
    return out_path


def main() -> int:
    for src, dest in DOCS:
        md_path = HERE / src
        if not md_path.is_file():
            raise SystemExit(f"missing source: {md_path}")
        out = render(md_path, HERE / dest)
        print(f"wrote {out.relative_to(HERE.parent)}  ({out.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
